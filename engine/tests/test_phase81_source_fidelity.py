from __future__ import annotations

import io
import zipfile
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw

from engine.src.memo_engine.canonical import (
    _attach_unassigned_media_context,
    build_canonical_memo,
)
from engine.src.memo_engine.ingestion import ingest_bytes
from engine.src.memo_engine.normalization import normalize_source
from engine.src.memo_engine.renderer import render_docx
from engine.src.memo_engine.structure import extract_structure, flatten_units


def _png(index: int, *, wide: bool = False) -> io.BytesIO:
    size = (320, 32) if wide else (180 + index, 120 + index)
    image = Image.new("RGB", size, (20 * index % 255, 80, 160))
    ImageDraw.Draw(image).text((8, 8), f"figure {index}", fill="white")
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    return stream


def _source_docx() -> bytes:
    document = Document()
    table = document.add_table(rows=0, cols=2)

    row = table.add_row()
    row.cells[0].paragraphs[0].add_run().add_picture(_png(1))

    row = table.add_row()
    row.cells[0].paragraphs[0].add_run("1.1 answer")

    row = table.add_row()
    p = row.cells[0].paragraphs[0]
    p.add_run("2.1 answer ")
    p.add_run().add_picture(_png(2))

    for major, image_index in ((3, 3), (4, 4)):
        row = table.add_row()
        row.cells[0].paragraphs[0].add_run().add_picture(_png(image_index))
        row.cells[0].add_paragraph(f"{major}.1 answer\n{major}.2 answer")

    row = table.add_row()
    row.cells[0].paragraphs[0].add_run().add_picture(_png(5, wide=True))
    row.cells[0].add_paragraph().add_run().add_picture(_png(6, wide=True))
    row.cells[0].add_paragraph("5.1.1 answer\n5.1.2 answer")
    row.cells[1].paragraphs[0].add_run().add_picture(_png(7))

    row = table.add_row()
    row.cells[0].paragraphs[0].add_run().add_picture(_png(8))
    row.cells[0].add_paragraph("6.2 answer\n6.3 answer")

    row = table.add_row()
    row.cells[0].paragraphs[0].add_run("9.1 answer")
    row.cells[0].add_paragraph().add_run().add_picture(_png(9))
    row.cells[1].add_paragraph().add_run().add_picture(_png(10, wide=True))

    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _pipeline(source: bytes):
    job = {
        "id": "phase81",
        "user_id": "teacher",
        "source_path": "teacher/phase81/source/input.docx",
        "source_filename": "input.docx",
        "source_mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }
    ingestion = ingest_bytes(job, source)
    normalized = normalize_source(job, source, ingestion)
    structure = extract_structure(normalized)
    structure["questions"].append({
        "question_id": "6.1",
        "path": [6, 1],
        "depth": 2,
        "context_only": False,
        "source_block_index": None,
        "source_preview": "Teacher-inserted missing subquestion",
        "printed_marks": None,
        "computed_shorthand_marks": None,
        "mark_points": [],
        "correction_overlay": {
            "correction_id": "correction-1",
            "operation": "insert_missing_child_question",
            "parent_id": "6",
        },
    })
    semantic = {
        "phase": "phase4_semantics",
        "deterministic_results": [],
        "ai_results": [],
        "exceptions": [],
        "interpreter_runs": [],
    }
    canonical, _, _ = build_canonical_memo(
        job, ingestion, normalized, structure, semantic, source
    )
    return normalized, structure, canonical


def _figures(canonical: dict) -> list[tuple[str, dict]]:
    result: list[tuple[str, dict]] = []

    def walk(items):
        for item in items:
            result.extend((str(item["number"]), block) for block in item.get("context_blocks", []) if block.get("type") == "figure")
            for alt in item.get("alternatives", []):
                result.extend((str(item["number"]), block) for block in alt.get("blocks", []) if block.get("type") == "figure")
            walk(item.get("children", []))

    for question in canonical["questions"]:
        result.extend((f"Q{question['number']}", block) for block in question.get("context_blocks", []) if block.get("type") == "figure")
        walk(question.get("items", []))
    return result


def test_normalization_promotes_media_and_keeps_image_only_rows():
    source = _source_docx()
    normalized, _, _ = _pipeline(source)
    entities = normalized["content"]["media_entities"]

    assert len(entities) == 10
    assert len(flatten_units(normalized)) == 8
    first = entities[0]
    assert first["relationship_id"].startswith("rId")
    assert first["package_path"].startswith("word/media/")
    assert first["sha256"]
    assert first["pixel_width"] and first["pixel_height"]
    assert first["raw_anchor"] == {
        "unit_index": 0,
        "row_index": 0,
        "cell_index": 0,
        "paragraph_index": 0,
        "drawing_order": 0,
    }
    assert first["source_block_index"] == 0
    assert any(entity["cell_role"] == "marking_allocation" for entity in entities)


def test_context_ownership_preserves_all_media_and_q6_dual_provenance():
    _, _, canonical = _pipeline(_source_docx())
    figures = _figures(canonical)
    assert len(figures) == 10
    assert len({block["figure"]["asset_id"] for _, block in figures}) == 10
    assert sum(block["figure"]["media_class"] == "substantive_figure" for _, block in figures) == 7
    owners = [owner for owner, _ in figures]
    assert "Q1" in owners
    assert "2.1" in owners
    assert "Q3" in owners
    assert "Q4" in owners
    assert owners.count("Q5") == 3
    assert "6.1" in owners
    assert "9.1" in owners
    q6 = next(block for owner, block in figures if owner == "6.1")
    assert q6["source_refs"][0]["docx_anchor"]["row_index"] == 6
    q6_item = next(
        item
        for question in canonical["questions"] if question["number"] == "6"
        for item in question["items"] if item["number"] == "6.1"
    )
    assert q6_item["correction_origin"]["source_block_index"] is None


def test_renderer_emits_seven_substantive_context_figures(tmp_path: Path, monkeypatch):
    source = _source_docx()
    _, _, canonical = _pipeline(source)
    substantive = {
        asset["asset_id"]
        for asset in canonical["source"]["assets"]
        if asset.get("media_class") == "substantive_figure"
    }
    for question in canonical["questions"]:
        question["context_blocks"] = [
            block for block in question.get("context_blocks", [])
            if (block.get("figure") or {}).get("asset_id") in substantive
        ]
        def retain_substantive(items):
            for item in items:
                item["context_blocks"] = [
                    block for block in item.get("context_blocks", [])
                    if (block.get("figure") or {}).get("asset_id") in substantive
                ]
                for alt in item.get("alternatives", []):
                    alt["blocks"] = [
                        block for block in alt.get("blocks", [])
                        if block.get("type") != "figure" or (block.get("figure") or {}).get("asset_id") in substantive
                    ]
                retain_substantive(item.get("children", []))

        retain_substantive(question.get("items", []))
    canonical["status"] = "render_ready"
    monkeypatch.setattr("engine.src.memo_engine.renderer._pandoc_math_bank", lambda _: {})
    output = tmp_path / "media.docx"
    render_docx(canonical, source, output, page_count=1)
    with zipfile.ZipFile(output) as archive:
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    assert len(media) == 7


def _corrected_child(question_id: str) -> dict:
    return {
        "question_id": question_id,
        "source_block_index": None,
        "correction_overlay": {
            "correction_id": f"correction-{question_id}",
            "operation": "insert_missing_child_question",
        },
    }


def test_question_12_inserted_child_inherits_only_unambiguous_preceding_media():
    items = [
        {"number": qid, "context_blocks": [], "children": [], "alternatives": []}
        for qid in ("12.1", "12.2", "12.3")
    ]
    questions = [{"number": "12", "context_blocks": [], "items": items}]
    anchor = {
        "unit_index": 2, "row_index": 4, "cell_index": 0,
        "paragraph_index": 0, "drawing_order": 0, "page": 3,
    }
    assets = [{
        "asset_id": "asset_portable",
        "media_class": "substantive_figure",
        "cell_role": "source",
        "source_anchor": anchor,
    }]
    records = [{
        "block_index": 8,
        "media_assets": ["asset_portable"],
        "paragraphs": [
            {
                "text": "",
                "source_ref": {"docx_anchor": {**anchor, "page": None}},
            },
            {
                "text": "12.2 answer\n12.3 answer",
                "source_ref": {"docx_anchor": {
                    "unit_index": 2, "row_index": 4, "cell_index": 0,
                    "paragraph_index": 1,
                }},
            },
        ],
    }]
    qmap = {
        "12.1": _corrected_child("12.1"),
        "12.2": {"question_id": "12.2", "source_block_index": 8},
        "12.3": {"question_id": "12.3", "source_block_index": 8},
    }

    _attach_unassigned_media_context(questions, records, qmap, assets)

    assert questions[0]["context_blocks"] == []
    assert len(items[0]["context_blocks"]) == 1
    block = items[0]["context_blocks"][0]
    assert block["figure"]["asset_id"] == "asset_portable"
    assert block["source_refs"][0]["docx_anchor"]["row_index"] == 4
    assert qmap["12.1"]["source_block_index"] is None


def test_ambiguous_inserted_predecessors_keep_media_at_parent_context():
    items = [
        {"number": qid, "context_blocks": [], "children": [], "alternatives": []}
        for qid in ("12.1", "12.2", "12.3")
    ]
    questions = [{"number": "12", "context_blocks": [], "items": items}]
    anchor = {
        "unit_index": 2, "row_index": 4, "cell_index": 0,
        "paragraph_index": 0, "drawing_order": 0, "page": 3,
    }
    assets = [{
        "asset_id": "asset_ambiguous",
        "media_class": "substantive_figure",
        "cell_role": "source",
        "source_anchor": anchor,
    }]
    records = [{
        "block_index": 8,
        "media_assets": ["asset_ambiguous"],
        "paragraphs": [{
            "text": "12.3 answer",
            "source_ref": {"docx_anchor": {
                "unit_index": 2, "row_index": 4, "cell_index": 0,
                "paragraph_index": 1,
            }},
        }],
    }]
    qmap = {
        "12.1": _corrected_child("12.1"),
        "12.2": _corrected_child("12.2"),
        "12.3": {"question_id": "12.3", "source_block_index": 8},
    }

    _attach_unassigned_media_context(questions, records, qmap, assets)

    assert len(questions[0]["context_blocks"]) == 1
    assert all(item["context_blocks"] == [] for item in items)
