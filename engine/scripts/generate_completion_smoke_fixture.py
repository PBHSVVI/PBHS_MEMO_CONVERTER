from __future__ import annotations

import argparse
import io
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT = ROOT / "tmp" / "completion_smoke_memo.docx"
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_FIXED_DOCUMENT_TIME = datetime(2026, 1, 1, tzinfo=timezone.utc)


def _set_cell(cell, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    lines = text.splitlines() or [""]
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        run = paragraph.add_run(line)
        run.bold = bold
        run.font.name = "Arial"
        run.font.size = Pt(10)


def _canonicalize_docx(data: bytes) -> bytes:
    source = io.BytesIO(data)
    target = io.BytesIO()
    with zipfile.ZipFile(source, "r") as archive, zipfile.ZipFile(
        target, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as output:
        for name in sorted(archive.namelist()):
            info = zipfile.ZipInfo(name, _FIXED_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            output.writestr(info, archive.read(name))
    return target.getvalue()


def build_fixture_bytes() -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    properties = document.core_properties
    properties.title = "PBHS deterministic completion smoke memo"
    properties.subject = "Mathematics lifecycle smoke fixture"
    properties.author = "PBHS Mathematics Memo Converter"
    properties.last_modified_by = "PBHS Mathematics Memo Converter"
    properties.created = _FIXED_DOCUMENT_TIME
    properties.modified = _FIXED_DOCUMENT_TIME
    properties.keywords = "public synthetic deterministic smoke fixture"
    properties.comments = "Contains no learner, teacher, or production memo data."

    header_lines = (
        "PRETORIA BOYS HIGH SCHOOL",
        "MATHEMATICS",
        "2026",
        "CLASS TEST",
        "PAPER 1",
        "FORM 5",
        "TIME: 0.5 HOURS (30 MINUTES)",
        "MARKS: 5",
    )
    for index, line in enumerate(header_lines):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line)
        run.bold = index in {0, 1, 3}
        run.font.name = "Arial"
        run.font.size = Pt(12 if index < 2 else 10)

    heading = document.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(4)
    run = heading.add_run("QUESTION 1")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, label in zip(
        table.rows[0].cells,
        ("QUESTION", "MEMO ANSWER", "MARKING GUIDELINE", "MARKS"),
    ):
        _set_cell(cell, label, bold=True)

    rows = (
        ("1.1", "The sum of 2 and 3 is 5.", "1A Correct answer", "(1)"),
        (
            "1.2\nSolve for x.",
            "Subtract 2 from both sides.\nx has the value 3.",
            "1M Correct method\n1A Correct answer",
            "(2)",
        ),
        (
            "1.3\nEvaluate the square of 2.",
            "Substitute 2 and square it.\nThe answer is 4.",
            "1M Correct substitution method\n1A Correct answer",
            "(2)",
        ),
    )
    for values in rows:
        row = table.add_row()
        for cell, value in zip(row.cells, values):
            _set_cell(cell, value)

    subtotal = table.add_row()
    _set_cell(subtotal.cells[0], "[5]", bold=True)
    for cell in subtotal.cells[1:]:
        _set_cell(cell, "")

    total = document.add_paragraph()
    total.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    total.paragraph_format.space_before = Pt(6)
    run = total.add_run("TOTAL: 5")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)

    raw = io.BytesIO()
    document.save(raw)
    return _canonicalize_docx(raw.getvalue())


def generate_fixture(path: str | Path = DEFAULT_OUTPUT) -> Path:
    output = Path(path).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_bytes(build_fixture_bytes())
    return output


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Generate the public deterministic hosted-completion smoke DOCX."
    )
    parser.add_argument("output", nargs="?", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    output = generate_fixture(args.output)
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
