from __future__ import annotations

import copy
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

RENDERER_VERSION = "phase6.0"
RENDER_PROFILE = "PBHS_GDE_INTERNAL_V1"

REQUIRED_GLYPHS = "∴≤≥≠±−×÷√∈∪∩∅∞∠⟂∥≅✓"


class RenderingError(RuntimeError):
    def __init__(
        self,
        code: str,
        message: str,
        *,
        details: dict[str, Any] | None = None,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.public_message = message
        self.details = details or {}


def _iter_items(items: list[dict[str, Any]]) -> Iterable[dict[str, Any]]:
    for item in items:
        yield item
        yield from _iter_items(item.get("children", []))


def _iter_blocks(canonical: dict[str, Any]) -> Iterable[dict[str, Any]]:
    for question in canonical.get("questions", []):
        for item in _iter_items(question.get("items", [])):
            for block in item.get("context_blocks", []):
                yield block
            for alternative in item.get("alternatives", []):
                for block in alternative.get("blocks", []):
                    yield block


def _set_cell_margins(cell, top=55, start=70, bottom=55, end=70) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_width(cell, twips: int) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(twips))
    tcW.set(qn("w:type"), "dxa")


def _set_table_fixed(table, widths_twips: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths_twips)))
    tblW.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_twips:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths_twips):
                _set_cell_width(cell, widths_twips[idx])
                _set_cell_margins(cell)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def _set_cell_borders(cell, size: int = 4) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")


def _set_row_cant_split(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    node = trPr.find(qn("w:cantSplit"))
    if node is None:
        node = OxmlElement("w:cantSplit")
        trPr.append(node)


def _set_row_keep_next(row) -> None:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            pPr = paragraph._p.get_or_add_pPr()
            keep = pPr.find(qn("w:keepNext"))
            if keep is None:
                keep = OxmlElement("w:keepNext")
                pPr.append(keep)


def _set_font(run, name: str = "Times New Roman", size: float = 9, *, bold=False, italic=False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rFonts.set(qn(f"w:{attr}"), name)


def _format_paragraph(p, *, align=None, before=0, after=0, line=1.0, keep_next=False) -> None:
    if align is not None:
        p.alignment = align
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if keep_next:
        pPr = p._p.get_or_add_pPr()
        keep = pPr.find(qn("w:keepNext"))
        if keep is None:
            keep = OxmlElement("w:keepNext")
            pPr.append(keep)


def _clear_cell(cell) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    _format_paragraph(p)


def _append_page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    _set_font(run, size=9)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])


def _source_anchor(block: dict[str, Any]) -> tuple[Any, ...] | None:
    refs = block.get("source_refs") or []
    if not refs:
        return None
    anchor = refs[0].get("docx_anchor") if isinstance(refs[0], dict) else None
    if not isinstance(anchor, dict):
        return None
    return (
        anchor.get("unit_index"), anchor.get("row_index"),
        anchor.get("cell_index"), anchor.get("paragraph_index"),
    )


def _group_blocks(blocks: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    result: list[list[dict[str, Any]]] = []
    for block in blocks:
        key = _source_anchor(block)
        if (
            result
            and key is not None
            and key == _source_anchor(result[-1][-1])
            and block.get("type") != "figure"
            and result[-1][-1].get("type") != "figure"
        ):
            result[-1].append(block)
        else:
            result.append([block])
    return result


def _prepare_latex(latex: str) -> str:
    """Keep mathematics native while forcing accidental prose inside OMML upright."""
    holders: list[str] = []

    def protect(match: re.Match[str]) -> str:
        holders.append(match.group(0))
        return f"@@P{len(holders) - 1}@@"

    value = re.sub(r"\\text\{[^{}]*\}", protect, latex)
    value = re.sub(r"\\[A-Za-z]+", protect, value)
    functions = {"sin", "cos", "tan", "log", "ln", "lim", "max", "min"}

    def word(match: re.Match[str]) -> str:
        token = match.group(0)
        if token.lower() in functions or token.isupper():
            return token
        return r"\text{" + token + "}"

    value = re.sub(r"\b[A-Za-z]{3,}\b", word, value)
    for index, original in enumerate(holders):
        value = value.replace(f"@@P{index}@@", original)

    # Normalise Unicode operators for Pandoc's TeX reader. The canonical
    # source remains unchanged; this is only the DOCX-native math transport.
    value = (
        value.replace("−", "-")
        .replace("∪", r"\cup{}")
        .replace("∩", r"\cap{}")
        .replace("∥", r"\parallel{}")
        .replace("∠", r"\angle{}")
        .replace("≅", r"\cong{}")
    )

    # LibreOffice can display a red parse marker for half-open interval
    # delimiters emitted as literal mismatched brackets. TeX left/right pairs
    # preserve the exact interval semantics without that renderer artefact.
    interval = re.compile(r"([\(\[])([^;\n]+);([^\)\]\n]+)([\)\]])")
    def interval_repl(match: re.Match[str]) -> str:
        whole = match.group(0)
        if r"\left" in whole or r"\right" in whole:
            return whole
        return (
            r"\left" + match.group(1)
            + match.group(2) + ";" + match.group(3)
            + r"\right" + match.group(4)
        )
    value = interval.sub(interval_repl, value)
    return value


def _math_display_parts(block: dict[str, Any]) -> list[tuple[str, str]]:
    math = block.get("math") or {}
    source = str(math.get("source_text") or "").strip()
    latex = str(math.get("canonical_latex") or "").strip()

    # Long English reasons embedded inside a Word equation do not wrap safely
    # in a narrow memo column. Split only the explanatory parenthetical; the
    # mathematical expression itself remains native OMML.
    m = re.match(r"^(.*?)(\s*\(([^()]*(?:[A-Za-z]{3,})[^()]*)\))\s*$", source)
    if m and m.group(1).strip():
        prefix_source = m.group(1).strip()
        prefix_latex = latex.split("(", 1)[0].strip()
        if prefix_latex:
            return [("math", _prepare_latex(prefix_latex)), ("text", " " + m.group(2).strip())]

    # Unit prose after a compact expression, e.g. n=3 years 2 months.
    m = re.match(
        r"^(.+?=\s*[0-9.,-]+)\s+((?:years?|months?|weeks?|days?|hours?|minutes?)\b.*)$",
        source,
        re.I,
    )
    if m:
        # The prefix is deliberately taken from the source because this narrow
        # unit pattern contains only a compact scalar expression before prose.
        prefix = m.group(1).strip()
        if prefix:
            return [("math", _prepare_latex(prefix)), ("text", " " + m.group(2))]

    prepared = _prepare_latex(latex)
    stripped = prepared.strip()
    if stripped == "=":
        return [("text", "=")]
    if stripped.startswith("="):
        rest = stripped[1:].strip()
        return [("text", "= ")] + ([("math", rest)] if rest else [])
    if stripped.endswith("="):
        rest = stripped[:-1].strip()
        return ([("math", rest)] if rest else []) + [("text", " =")]
    return [("math", prepared)]


def _all_math_variants(canonical: dict[str, Any]) -> list[str]:
    values: list[str] = []
    seen: set[str] = set()
    for block in _iter_blocks(canonical):
        if block.get("type") != "math":
            continue
        for kind, value in _math_display_parts(block):
            if kind == "math" and value and value not in seen:
                seen.add(value)
                values.append(value)
    return values


def _pandoc_math_bank(canonical: dict[str, Any]) -> dict[str, ET.Element]:
    latex_values = _all_math_variants(canonical)
    if not latex_values:
        return {}
    pandoc = shutil.which("pandoc")
    if not pandoc:
        raise RenderingError(
            "RENDER_PANDOC_MISSING",
            "Pandoc is required to create native Word mathematics on the hosted renderer.",
        )

    with tempfile.TemporaryDirectory(prefix="pbhs-math-") as td:
        td_path = Path(td)
        md_path = td_path / "math.md"
        docx_path = td_path / "math.docx"
        md_path.write_text(
            "\n\n".join(f"$${value}$$" for value in latex_values),
            encoding="utf-8",
        )
        result = subprocess.run(
            [pandoc, str(md_path), "-o", str(docx_path)],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=False,
            timeout=120,
        )
        if result.returncode != 0:
            raise RenderingError(
                "RENDER_MATH_CONVERSION_FAILED",
                "Pandoc could not convert canonical mathematics to native Word equations.",
            )
        with zipfile.ZipFile(docx_path) as archive:
            root = ET.fromstring(archive.read("word/document.xml"))
        equations = root.findall(f".//{{{M_NS}}}oMath")
        if len(equations) != len(latex_values):
            raise RenderingError(
                "RENDER_MATH_COUNT_MISMATCH",
                "Native equation conversion did not preserve one equation per canonical expression.",
            )
        return {latex: equation for latex, equation in zip(latex_values, equations)}


def _append_omml(paragraph, latex: str, bank: dict[str, ET.Element]) -> None:
    equation = bank.get(latex)
    if equation is None:
        raise RenderingError(
            "RENDER_MATH_LOOKUP_MISSING",
            "A canonical equation was missing from the native Word math bank.",
        )
    element = parse_xml(ET.tostring(equation, encoding="unicode"))
    # The approved PBHS gold files use body text at 9 pt with mathematical
    # expressions visually one step larger. Explicit 10 pt math also avoids
    # Linux font-substitution compressing equations too aggressively.
    for rPr in element.iter(qn("w:rPr")):
        for tag in ("w:sz", "w:szCs"):
            node = rPr.find(qn(tag))
            if node is None:
                node = OxmlElement(tag)
                rPr.append(node)
            node.set(qn("w:val"), "20")
    paragraph._p.append(element)


def _add_text(paragraph, text: str, *, bold=False, italic=False, size=9) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    _set_font(run, size=size, bold=bold, italic=italic)


def _asset_bytes(source_bytes: bytes, canonical: dict[str, Any]) -> dict[str, tuple[bytes, str]]:
    assets = {item["asset_id"]: item for item in canonical.get("source", {}).get("assets", [])}
    result: dict[str, tuple[bytes, str]] = {}
    if not assets:
        return result
    try:
        with zipfile.ZipFile(io.BytesIO(source_bytes)) as archive:
            for asset_id, meta in assets.items():
                path = meta.get("package_path")
                if not path:
                    continue
                try:
                    result[asset_id] = (archive.read(path), str(meta.get("mime_type") or "application/octet-stream"))
                except KeyError:
                    continue
    except zipfile.BadZipFile as exc:
        raise RenderingError("RENDER_SOURCE_PACKAGE_INVALID", "The source DOCX could not be reopened for figure rendering.") from exc
    return result


def _add_figure(paragraph, block: dict[str, Any], assets: dict[str, tuple[bytes, str]]) -> None:
    asset_id = (block.get("figure") or {}).get("asset_id")
    blob = assets.get(str(asset_id))
    if blob is None:
        raise RenderingError("RENDER_ASSET_MISSING", "A canonical figure asset could not be resolved during rendering.")
    data, mime = blob
    suffix = ".png"
    if "jpeg" in mime:
        suffix = ".jpg"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as temp:
        temp.write(data)
        temp_path = temp.name
    try:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(temp_path, width=Inches(3.35))
    finally:
        try:
            os.unlink(temp_path)
        except OSError:
            pass


def _render_block_group(cell, group: list[dict[str, Any]], bank, assets) -> None:
    if len(group) == 1 and group[0].get("type") == "figure":
        p = cell.add_paragraph() if cell.paragraphs[0].text or len(cell.paragraphs[0]._p) > 1 else cell.paragraphs[0]
        _format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=1)
        _add_figure(p, group[0], assets)
        return

    p = cell.add_paragraph() if cell.paragraphs[0].text or len(cell.paragraphs[0]._p) > 1 else cell.paragraphs[0]
    _format_paragraph(p, after=0)
    contains_display_math = False
    previous_kind: str | None = None
    for block in group:
        kind = block.get("type")
        if previous_kind is not None and (kind == "prose" or previous_kind == "prose"):
            _add_text(p, " ")
        if kind == "prose":
            _add_text(p, str(block.get("text") or ""))
        elif kind == "math":
            parts = _math_display_parts(block)
            for part_index, (part_kind, value) in enumerate(parts):
                if part_index > 0 and part_kind == "text":
                    # _math_display_parts supplies its own leading space where needed.
                    pass
                if part_kind == "math":
                    _append_omml(p, value, bank)
                    contains_display_math = True
                else:
                    _add_text(p, value)
        elif kind == "note":
            _add_text(p, str(block.get("text") or ""))
        previous_kind = kind
    if contains_display_math and len(group) == 1 and group[0].get("type") == "math":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _descriptor(mark: dict[str, Any]) -> str:
    value = str(mark.get("descriptor") or "").strip()
    mark_type = str(mark.get("type") or "")
    if mark_type == "consistent_accuracy" and "(CA)" not in value.upper():
        value = (value or "answer") + " (CA)"
    if mark_type == "reason" and value and not value.lower().startswith("reason"):
        value = "reason: " + value
    return value or mark_type.replace("_", " ")


def _mark_signature(mark: dict[str, Any]) -> tuple[Any, ...]:
    ca = mark.get("consistent_accuracy") or {}
    return (
        str(mark.get("type") or ""),
        int(mark.get("count") or 0),
        str(mark.get("descriptor") or "").strip(),
        bool(ca.get("enabled")),
        tuple(str(x) for x in ca.get("dependency_item_ids", []) or []),
        tuple(str(x) for x in mark.get("acceptable_variants", []) or []),
    )


def _same_marking_scheme(alternatives: list[dict[str, Any]]) -> bool:
    if len(alternatives) < 2:
        return False
    signatures = [
        tuple(_mark_signature(mark) for mark in alt.get("marking_points", []))
        for alt in alternatives
    ]
    return bool(signatures[0]) and all(sig == signatures[0] for sig in signatures[1:])


def _render_marks(cell, alternatives: list[dict[str, Any]]) -> None:
    _clear_cell(cell)
    render_alternatives = alternatives[:1] if _same_marking_scheme(alternatives) else alternatives
    first = True
    for alt_index, alt in enumerate(render_alternatives):
        if alt_index > 0:
            p = cell.add_paragraph() if not first else cell.paragraphs[0]
            _format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER)
            _add_text(p, "OR", bold=True)
            first = False
        for mark in alt.get("marking_points", []):
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            _format_paragraph(p, after=0)
            count = int(mark.get("count") or 0)
            tick = p.add_run("✓" * count)
            _set_font(tick, name="Segoe UI Symbol", size=9)
            _add_text(p, " " + _descriptor(mark))


def _render_working(cell, alternatives: list[dict[str, Any]], bank, assets) -> None:
    _clear_cell(cell)
    first_alt = True
    for alt_index, alt in enumerate(alternatives):
        if alt_index > 0:
            p = cell.add_paragraph()
            _format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=1, after=1)
            _add_text(p, "OR", bold=True)
        groups = _group_blocks(alt.get("blocks", []))
        for group in groups:
            _render_block_group(cell, group, bank, assets)
        first_alt = False


def _render_context(cell, blocks: list[dict[str, Any]], bank, assets) -> None:
    _clear_cell(cell)
    for group in _group_blocks(blocks):
        _render_block_group(cell, group, bank, assets)


def _leaf_row_should_avoid_split(item: dict[str, Any]) -> bool:
    alternatives = item.get("alternatives", [])
    blocks = [block for alt in alternatives for block in alt.get("blocks", [])]
    block_count = len(blocks)
    char_count = 0
    figure_count = 0
    for block in blocks:
        if block.get("type") == "figure":
            figure_count += 1
        elif block.get("type") == "math":
            math = block.get("math") or {}
            char_count += len(str(math.get("plain_text") or math.get("source_text") or ""))
        else:
            char_count += len(str(block.get("text") or ""))

    # Prevent Word from splitting ordinary memo rows independently by cell,
    # which can strand the question number/marks on one page and the working
    # on the next. Very large future rows remain splittable by design.
    return block_count <= 24 and char_count <= 450 and figure_count <= 1


def _add_item_rows(table, item: dict[str, Any], bank, assets) -> None:
    children = item.get("children", [])
    context_blocks = item.get("context_blocks", [])
    alternatives = item.get("alternatives", [])

    if children and context_blocks:
        row = table.add_row()
        for cell in row.cells:
            _set_cell_borders(cell)
            _set_cell_margins(cell)
        _clear_cell(row.cells[0])
        _add_text(row.cells[0].paragraphs[0], str(item.get("number") or ""))
        _render_context(row.cells[1], context_blocks, bank, assets)
        _clear_cell(row.cells[2]); _clear_cell(row.cells[3])

    if not children:
        row = table.add_row()
        if _leaf_row_should_avoid_split(item):
            _set_row_cant_split(row)
        for cell in row.cells:
            _set_cell_borders(cell)
            _set_cell_margins(cell)
        _clear_cell(row.cells[0])
        _add_text(row.cells[0].paragraphs[0], str(item.get("number") or ""))
        _render_working(row.cells[1], alternatives, bank, assets)
        _render_marks(row.cells[2], alternatives)
        _clear_cell(row.cells[3])
        p = row.cells[3].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _add_text(p, f"({int((item.get('marks') or {}).get('computed') or 0)})")

    for child in children:
        _add_item_rows(table, child, bank, assets)


def _cover(document: Document, canonical: dict[str, Any], page_count: int | None) -> None:
    meta = canonical.get("document_metadata", {})
    paper = str(meta.get("paper") or "PAPER 1")
    year = str(meta.get("year") or "")
    grade = str(meta.get("grade_label") or "")
    marks = canonical.get("totals", {}).get("computed")
    duration = meta.get("duration_minutes")

    for text, size in [
        (str(meta.get("exam_type") or "PREPARATORY EXAMINATION"), 13),
        (year, 13),
        ("MARKING GUIDELINES", 13),
        (f"MATHEMATICS {paper}", 13),
        (grade, 11),
    ]:
        p = document.add_paragraph()
        _format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
        _add_text(p, text.upper(), bold=True, size=size)

    for _ in range(8):
        document.add_paragraph()

    p = document.add_paragraph(); _format_paragraph(p, after=1)
    _add_text(p, f"MARKS: {marks}", bold=True, size=10)
    p = document.add_paragraph(); _format_paragraph(p, after=1)
    if duration:
        hours = float(duration) / 60.0
        label = str(int(hours)) if hours.is_integer() else str(hours).rstrip("0").rstrip(".")
        _add_text(p, f"TIME: {label} HOURS", bold=True, size=10)
    else:
        _add_text(p, "TIME:", bold=True, size=10)
    p = document.add_paragraph(); _format_paragraph(p, after=0)
    extent = str(page_count) if page_count else "__PAGECOUNT__"
    _add_text(p, f"This marking guideline consists of {extent} pages including the cover page.", bold=True, size=10)


def _notes(document: Document) -> None:
    document.add_page_break()
    p = document.add_paragraph(); _format_paragraph(p, after=5)
    _add_text(p, "NOTES:", bold=True, italic=True, size=10.5)
    notes = [
        "If a candidate answered a question TWICE, mark the FIRST attempt.",
        "If a candidate crossed out an answer and did not redo the question, mark the crossed-out answer.",
        "Consistent Accuracy (CA) applies to ALL aspects of the marking guideline.",
        "It is UNACCEPTABLE to assume values or answers in order to complete a solution.",
        "(A) denotes an accuracy mark.",
    ]
    for idx, note in enumerate(notes, 1):
        p = document.add_paragraph()
        _format_paragraph(p, after=4)
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        _add_text(p, f"{idx}. {note}", size=10)


def _question_table(document: Document, question: dict[str, Any], bank, assets, *, last=False, total=None) -> None:
    table = document.add_table(rows=0, cols=4)
    widths = [892, 6552, 2232, 648]
    _set_table_fixed(table, widths)

    heading = table.add_row()
    _set_row_cant_split(heading)
    _set_row_keep_next(heading)
    merged = heading.cells[0].merge(heading.cells[1]).merge(heading.cells[2]).merge(heading.cells[3])
    _set_cell_borders(merged)
    _clear_cell(merged)
    p = merged.paragraphs[0]
    _format_paragraph(p, keep_next=True)
    _add_text(p, str(question.get("heading") or f"QUESTION {question.get('number')}"), bold=True)

    for item in question.get("items", []):
        _add_item_rows(table, item, bank, assets)

    subtotal = table.add_row()
    _set_row_cant_split(subtotal)
    left = subtotal.cells[0].merge(subtotal.cells[1]).merge(subtotal.cells[2])
    for cell in (left, subtotal.cells[3]):
        _set_cell_borders(cell); _set_cell_margins(cell)
    _clear_cell(left); _clear_cell(subtotal.cells[3])
    p = subtotal.cells[3].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_text(p, f"[{int((question.get('subtotal') or {}).get('computed') or 0)}]", bold=True)

    if last:
        # Keep the final question subtotal with the TOTAL row so TOTAL cannot
        # become an otherwise-empty trailing page.
        _set_row_keep_next(subtotal)
        row = table.add_row()
        _set_row_cant_split(row)
        left = row.cells[0].merge(row.cells[1]).merge(row.cells[2])
        for cell in (left, row.cells[3]):
            _set_cell_borders(cell); _set_cell_margins(cell)
        _clear_cell(left); _clear_cell(row.cells[3])
        p = left.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _add_text(p, "TOTAL:", bold=True)
        p = row.cells[3].paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _add_text(p, str(int(total or 0)), bold=True)

    # Reapply fixed widths after row creation.
    _set_table_fixed(table, widths)


def _configure_section(document: Document, canonical: dict[str, Any]) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)
    section.different_first_page_header_footer = True

    meta = canonical.get("document_metadata", {})
    paper = str(meta.get("paper") or "PAPER 1").replace("PAPER ", "P")
    year = str(meta.get("year") or "2026")[-2:]

    header = section.header
    p = header.paragraphs[0]
    _format_paragraph(p, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _add_text(p, f"MATHEMATICS {paper}/{year}", bold=True, size=9)

    footer = section.footer
    _append_page_field(footer.paragraphs[0])
    first_footer = section.first_page_footer
    _append_page_field(first_footer.paragraphs[0])


def _set_doc_defaults(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(9)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")


def render_docx(
    canonical: dict[str, Any],
    source_bytes: bytes,
    output_path: str | Path,
    *,
    page_count: int | None = None,
) -> dict[str, Any]:
    if canonical.get("status") != "render_ready":
        raise RenderingError("RENDER_INPUT_NOT_READY", "The canonical memo is not render-ready.")
    if canonical.get("render_profile") != RENDER_PROFILE:
        raise RenderingError("RENDER_PROFILE_UNSUPPORTED", "The selected renderer profile is not supported by Phase 6.")

    bank = _pandoc_math_bank(canonical)
    assets = _asset_bytes(source_bytes, canonical)

    document = Document()
    _set_doc_defaults(document)
    _configure_section(document, canonical)
    _cover(document, canonical, page_count)
    _notes(document)
    document.add_page_break()

    questions = canonical.get("questions", [])
    total = canonical.get("totals", {}).get("computed")
    for index, question in enumerate(questions):
        _question_table(
            document,
            question,
            bank,
            assets,
            last=index == len(questions) - 1,
            total=total,
        )
        if index != len(questions) - 1:
            # A tiny zero-spacing paragraph separates native Word tables without
            # forcing a page break or visible blank line.
            p = document.add_paragraph()
            _format_paragraph(p, after=0, line=0.7)
            r = p.add_run("")
            _set_font(r, size=1)

    # Internal audit metadata; not printed.
    core = document.core_properties
    core.title = "PBHS Mathematics GDE Marking Guidelines"
    core.subject = f"Schema {canonical.get('schema_version')} / Renderer {RENDERER_VERSION} / {RENDER_PROFILE}"
    core.comments = f"document_id={canonical.get('document_id')}"

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return preflight_docx(output, canonical)


def preflight_docx(path: str | Path, canonical: dict[str, Any]) -> dict[str, Any]:
    path = Path(path)
    issues: list[str] = []
    try:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", errors="replace")
            settings = archive.read("word/settings.xml").decode("utf-8", errors="replace")
    except Exception as exc:
        raise RenderingError("RENDER_DOCX_INVALID", "The generated DOCX package could not be inspected.") from exc

    expected_math = sum(
        1
        for block in _iter_blocks(canonical)
        if block.get("type") == "math"
        for kind, _ in _math_display_parts(block)
        if kind == "math"
    )
    actual_math = xml.count("<m:oMath")
    if actual_math < expected_math:
        issues.append(f"MATH_COUNT:{actual_math}<{expected_math}")
    if "🗸" in xml:
        issues.append("LEGACY_TICK")
    if "�" in xml:
        issues.append("REPLACEMENT_CHARACTER")
    if "w:pgSz" not in xml or "w:pgMar" not in xml:
        issues.append("PAGE_SETTINGS_MISSING")
    if canonical.get("totals", {}).get("computed") == 150 and "TOTAL:" not in xml:
        issues.append("TOTAL_ROW_MISSING")
    if "✓" not in xml:
        issues.append("CHECK_MARK_MISSING")

    return {
        "passed": not issues,
        "issues": issues,
        "expected_math_blocks": expected_math,
        "native_math_elements": actual_math,
        "renderer_version": RENDERER_VERSION,
        "render_profile": RENDER_PROFILE,
    }


def convert_docx_to_pdf(docx_path: str | Path, output_dir: str | Path) -> Path:
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not libreoffice:
        raise RenderingError("RENDER_OFFICE_ENGINE_MISSING", "LibreOffice is required for deterministic PDF export.")
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    env = dict(os.environ)
    with tempfile.TemporaryDirectory(prefix="pbhs-lo-") as profile:
        env["HOME"] = profile
        result = subprocess.run(
            [
                libreoffice,
                "--headless",
                f"-env:UserInstallation=file://{profile}/profile",
                "--convert-to", "pdf",
                "--outdir", str(output_dir),
                str(docx_path),
            ],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            env=env,
            check=False,
            timeout=180,
        )
    pdf = output_dir / (Path(docx_path).stem + ".pdf")
    if result.returncode != 0 or not pdf.exists() or pdf.stat().st_size == 0:
        raise RenderingError("RENDER_PDF_EXPORT_FAILED", "The controlled office engine could not export the generated DOCX to PDF.")
    return pdf


def pdf_page_count(pdf_path: str | Path) -> int:
    try:
        from pypdf import PdfReader
        return len(PdfReader(str(pdf_path)).pages)
    except Exception as exc:
        raise RenderingError("RENDER_PDF_PAGECOUNT_FAILED", "The rendered PDF page count could not be determined.") from exc


def _memo_required_glyphs(canonical: dict[str, Any]) -> list[str]:
    visible: list[str] = []
    for block in _iter_blocks(canonical):
        if block.get("type") in {"prose", "note", "label"}:
            visible.append(str(block.get("text") or ""))
        elif block.get("type") == "math":
            math = block.get("math") or {}
            visible.append(str(math.get("source_text") or ""))
            visible.append(str(math.get("plain_text") or ""))
    text = "\n".join(visible)
    glyphs = [glyph for glyph in REQUIRED_GLYPHS if glyph in text]
    # Every scored memo renders check marks even when the source did not contain them.
    if any(
        alt.get("marking_points")
        for question in canonical.get("questions", [])
        for item in _iter_items(question.get("items", []))
        for alt in item.get("alternatives", [])
    ) and "✓" not in glyphs:
        glyphs.append("✓")
    return glyphs


def preflight_pdf(pdf_path: str | Path, canonical: dict[str, Any]) -> dict[str, Any]:
    from pypdf import PdfReader

    path = Path(pdf_path)
    issues: list[str] = []

    try:
        reader = PdfReader(str(path))
        page_text = [(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:
        raise RenderingError(
            "RENDER_PDF_INVALID",
            "The generated PDF could not be inspected.",
        ) from exc

    pypdf_text = "\n".join(page_text)

    # Poppler gives us an independent Unicode extraction path. A required
    # codepoint is considered preserved when at least one of the two standard
    # extractors can recover it. This avoids treating an extractor-specific
    # mapping quirk as a visible glyph failure while retaining a real codepoint
    # check as required by the frozen rendering specification.
    poppler_text = ""
    pdftotext = shutil.which("pdftotext")
    if pdftotext:
        result = subprocess.run(
            [pdftotext, "-enc", "UTF-8", str(path), "-"],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=False,
            timeout=30,
        )
        if result.returncode == 0:
            poppler_text = result.stdout.decode("utf-8", errors="replace")

    extraction_texts = [text for text in [pypdf_text, poppler_text] if text]
    primary_text = poppler_text or pypdf_text

    if not page_text:
        issues.append("PDF_NO_PAGES")
    if page_text and not page_text[-1].strip():
        issues.append("TRAILING_BLANK_PAGE")

    # A replacement character must be absent from every successful extraction.
    if any("�" in text for text in extraction_texts):
        issues.append("PDF_REPLACEMENT_CHARACTER")

    if "TOTAL:" not in primary_text:
        issues.append("PDF_TOTAL_MISSING")
    expected_total = canonical.get("totals", {}).get("computed")
    if expected_total is not None and str(int(expected_total)) not in primary_text:
        issues.append("PDF_TOTAL_VALUE_MISSING")

    required = _memo_required_glyphs(canonical)
    missing = [
        glyph
        for glyph in required
        if not any(glyph in text for text in extraction_texts)
    ]
    issues.extend(f"PDF_GLYPH_MISSING:{glyph}" for glyph in missing)

    fonts: list[dict[str, Any]] = []
    pdffonts = shutil.which("pdffonts")
    if pdffonts:
        result = subprocess.run(
            [pdffonts, str(path)],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=False,
            timeout=30,
            text=True,
        )
        if result.returncode == 0:
            lines = result.stdout.splitlines()[2:]
            for line in lines:
                parts = line.split()
                if len(parts) < 8:
                    continue

                # pdffonts always ends rows with:
                # encoding  emb  sub  uni  object-ID  generation.
                # Parse relative to the row end so variable-width / multi-word
                # font type labels cannot shift the embedding column.
                encoding = parts[-6]
                embedded = parts[-5].lower() == "yes"
                font_type = " ".join(parts[1:-6])
                fonts.append({
                    "name": parts[0],
                    "type": font_type,
                    "encoding": encoding,
                    "embedded": embedded,
                })

            if fonts and any(not item["embedded"] for item in fonts):
                issues.append("PDF_FONT_NOT_EMBEDDED")

    result = {
        "passed": not issues,
        "issues": issues,
        "page_count": len(page_text),
        "required_glyphs": required,
        "missing_glyphs": missing,
        "fonts": fonts,
        "text_extraction": {
            "pypdf_available": True,
            "pypdf_chars": len(pypdf_text),
            "pdftotext_available": bool(pdftotext),
            "pdftotext_chars": len(poppler_text),
            "total_found_pypdf": "TOTAL:" in pypdf_text,
            "total_found_pdftotext": "TOTAL:" in poppler_text if poppler_text else None,
        },
    }

    if issues:
        print(
            "Phase 6 PDF preflight failed: "
            + json.dumps(result, ensure_ascii=False, separators=(",", ":"))
        )

    return result


def render_outputs(
    canonical: dict[str, Any],
    source_bytes: bytes,
    output_dir: str | Path,
) -> dict[str, Any]:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    docx_path = output_dir / "memo.docx"

    # First pass obtains actual pagination from the same office engine used for
    # the final PDF. Second pass writes the cover extent and exports again.
    draft_preflight = render_docx(canonical, source_bytes, docx_path, page_count=None)
    if not draft_preflight["passed"]:
        raise RenderingError("RENDER_DOCX_PREFLIGHT_FAILED", "The generated DOCX failed structural preflight.")
    draft_pdf = convert_docx_to_pdf(docx_path, output_dir)
    pages = pdf_page_count(draft_pdf)

    final_preflight = render_docx(canonical, source_bytes, docx_path, page_count=pages)
    if not final_preflight["passed"]:
        raise RenderingError("RENDER_DOCX_PREFLIGHT_FAILED", "The final DOCX failed structural preflight.")
    pdf_path = convert_docx_to_pdf(docx_path, output_dir)
    final_pages = pdf_page_count(pdf_path)
    if final_pages != pages:
        # One bounded correction pass if the extent text itself changed pagination.
        pages = final_pages
        final_preflight = render_docx(canonical, source_bytes, docx_path, page_count=pages)
        pdf_path = convert_docx_to_pdf(docx_path, output_dir)
        final_pages = pdf_page_count(pdf_path)
    if final_pages != pages:
        raise RenderingError("RENDER_PAGECOUNT_UNSTABLE", "DOCX/PDF pagination did not stabilise after the extent update.")

    pdf_preflight = preflight_pdf(pdf_path, canonical)
    if not pdf_preflight["passed"]:
        issue_text = ", ".join(pdf_preflight.get("issues", [])[:8]) or "unknown"
        raise RenderingError(
            "RENDER_PDF_PREFLIGHT_FAILED",
            "The generated PDF failed deterministic preflight: " + issue_text,
            details=pdf_preflight,
        )

    return {
        "docx_path": str(docx_path),
        "pdf_path": str(pdf_path),
        "page_count": final_pages,
        "docx_preflight": final_preflight,
        "pdf_preflight": pdf_preflight,
    }
